"""Compare rules extracted from code against a standard rulebook (JSON).

Each standard rule gets one status:
  Matched            - the code implements the rule with the same values and outcome
  Different          - the code implements the rule, but a value or outcome differs
  Missing in code    - no code rule implements it
Every extracted business rule that no standard rule accounts for is reported as
  Additional in code - behaviour in the code that the standard does not describe

Matching is on structured facts (field keyword, operator, value), not wording, so it
works the same whichever backend (heuristic, anthropic, qwen) produced the text.
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

MATCHED, DIFFERENT, MISSING, ADDITIONAL = "Matched", "Different", "Missing in code", "Additional in code"
BUSINESS_TYPES = {"Validation", "Adjustment", "Calculation", "Condition", "Decision"}


# ---------------------------------------------------------------- helpers
def load_standard(path: str | Path) -> dict:
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    if "rules" not in data or not isinstance(data["rules"], list):
        raise ValueError("standard JSON must contain a 'rules' list")
    return data


def _kw(x) -> list[str]:
    return [k.upper() for k in (x if isinstance(x, list) else [x])]


def _field_hit(keywords, fields: list[str]) -> bool:
    """A keyword matches a field if it equals the field or one of its hyphen-separated parts."""
    for f in fields:
        f = (f or "").upper()
        parts = set(f.split("-"))
        if any(k == f or k in parts for k in _kw(keywords)):
            return True
    return False


def _all_parts(keywords, field: str) -> bool:
    parts = set((field or "").upper().split("-"))
    return all(k in parts for k in _kw(keywords))


def _eq(a, b) -> bool:
    try:
        return abs(float(a) - float(b)) < 1e-6
    except (TypeError, ValueError):
        return str(a).strip().upper() == str(b).strip().upper()


def _fmt(v) -> str:
    if isinstance(v, list):
        return ", ".join(_fmt(x) for x in v)
    if isinstance(v, float):
        return f"{v:,.2f}" if v >= 1 or v == 0 else f"{v:g}"
    return str(v)


_OPS = {"=": "is", "!=": "is not", ">": "above", ">=": "at least", "<": "below", "<=": "at most",
        "not_in": "not one of", "none_of": "none of", "one_of": "one of", "is": "is"}


def _cond_text(fields, op, value) -> str:
    return f"{'/'.join(_kw(fields))} {_OPS.get(op, op)} {_fmt(value)}"


# ---------------------------------------------------------------- code items
def _code_items(programs: list[dict]) -> list[dict]:
    """Flatten extracted rules into comparable items (table rows become their own items)."""
    items = []
    for p in programs:
        for r in p["rules"]:
            f = r.get("facts") or {}
            base = {"program": p["program"], "rule_id": r["id"], "type": r["type"], "rule": r["rule"],
                    "lines": f'{p["file"]}:{r["lines"][0]}' + (f'-{r["lines"][1]}' if r["lines"][1] != r["lines"][0] else "")}
            if f.get("kind") == "table":
                for row in f.get("rows", []):
                    if row["key"] == "OTHER":
                        continue
                    acts = "; ".join(f'{a.get("target", "")} = {_fmt(a.get("value", a.get("expr", "")))}'
                                     for a in row["actions"] if a["type"] in ("set", "calculate"))
                    items.append({**base, "key": f'{r["id"]}:{row["key"]}', "facts": {**f, "kind": "row", "row": row},
                                  "rule": f'When the {"/".join(f.get("fields", []))} is {row["key"]}: {acts or "no action"}',
                                  "lines": f'{p["file"]}:{row["line"]}'})
                items.append({**base, "key": f'{r["id"]}:table', "facts": f, "internal": True})
            else:
                items.append({**base, "key": r["id"], "facts": f,
                              "internal": r["type"] not in BUSINESS_TYPES})
    return items


# ---------------------------------------------------------------- matchers
def _match_condition(std: dict, items: list[dict]):
    w = std["when"]
    same_field_op, same_field = [], []
    for it in items:
        f = it["facts"]
        if f.get("kind") != "condition":
            continue
        for c in f.get("conditions", []):
            if not _field_hit(w["field"], c.get("fields") or [c.get("field", "")]):
                continue
            (same_field_op if c["op"] == w["op"] else same_field).append((it, c))
    scored = []
    for it, c in same_field_op:
        want, got = w.get("value"), c["value"]
        if isinstance(want, list):
            ws, gs = {str(x).upper() for x in want}, {str(x).upper() for x in (got if isinstance(got, list) else [got])}
            value_ok = ws == gs
            diff = []
            if ws - gs:
                diff.append(f"missing in code: {', '.join(sorted(ws - gs))}")
            if gs - ws:
                diff.append(f"extra in code: {', '.join(sorted(gs - ws))}")
            value_note = "; ".join(diff)
        else:
            value_ok = want is None or _eq(want, got)
            value_note = "" if value_ok else f"standard {_fmt(want)}, code {_fmt(got)}"
        outcome_ok, outcome_note = _check_then(std.get("then", []), it["facts"])
        code_has = f"{_cond_text(c.get('field'), c['op'], got)} → {_describe_actions(it['facts'].get('actions', []))}"
        # best = value and outcome agree; then outcome agrees; then value agrees; unused before used
        score = (value_ok and outcome_ok, outcome_ok, value_ok, it["key"] not in _USED)
        status = MATCHED if value_ok and outcome_ok else DIFFERENT
        scored.append((score, status, it, code_has, "; ".join(x for x in (value_note, outcome_note) if x)))
    if scored:
        scored.sort(key=lambda x: x[0], reverse=True)
        _, status, it, code_has, detail = scored[0]
        return status, it, code_has, detail
    return MISSING, None, "", ""


def _check_then(expect: list[dict], facts: dict) -> tuple[bool, str]:
    acts = facts.get("actions", [])
    problems = []
    for e in expect:
        if "set" in e:
            hits = [a for a in acts if a["type"] == "set" and _field_hit(e["set"], [a["target"]])]
            if not hits:
                problems.append(f"no {'/'.join(_kw(e['set']))} set in code")
            elif "value" in e and not any(_eq(e["value"], a["value"]) for a in hits):
                problems.append(f"{'/'.join(_kw(e['set']))}: standard {_fmt(e['value'])}, "
                                f"code {', '.join(_fmt(a['value']) for a in hits)}")
        if "calculate" in e:
            hits = [a for a in acts if a["type"] == "calculate" and _field_hit(e["calculate"], [a["target"]])]
            if not hits:
                problems.append(f"no {'/'.join(_kw(e['calculate']))} calculation in code")
            elif "factor" in e and not any(any(_eq(e["factor"], k) for k in a.get("constants", [])) for a in hits):
                problems.append(f"factor: standard {e['factor']}, code {hits[0]['expr']}")
    return (not problems), "; ".join(problems)


def _describe_actions(acts: list[dict]) -> str:
    out = []
    for a in acts:
        if a["type"] == "set":
            out.append(f'{a["target"]} = {_fmt(a["value"])}')
        elif a["type"] == "calculate":
            out.append(f'{a["target"]} = {a["expr"]}')
        elif a["type"] == "message":
            out.append(f'message "{a["value"]}"')
        elif a["type"] in ("perform", "goto", "call"):
            out.append(f'{a["type"]} {a["target"]}')
        elif a["type"] == "stop":
            out.append("stop")
    return "; ".join(out) or "no action"


def _tables(items):
    return [it for it in items if it["facts"].get("kind") == "table"]


def _match_lookup(std: dict, items: list[dict]):
    key = std["key"].upper()
    for it in items:
        f = it["facts"]
        if f.get("kind") != "row" or not _field_hit(std["table"]["field"], f.get("fields", [])):
            continue
        if f["row"]["key"].upper() != key:
            continue
        hits = [a for a in f["row"]["actions"] if a["type"] == "set" and _field_hit(std["sets"], [a["target"]])]
        pct = lambda v, t: f"{_fmt(v)} ({float(v) * 100:.2f}%)" if "PCT" in t.upper() and isinstance(v, float) else _fmt(v)
        code_has = "; ".join(f'{a["target"]} = {pct(a["value"], a["target"])}' for a in f["row"]["actions"]) or "no action"
        if hits and _eq(std["value"], hits[0]["value"]):
            return MATCHED, it, code_has, ""
        tgt = hits[0]["target"] if hits else "".join(_kw(std["sets"]))
        got = pct(hits[0]["value"], tgt) if hits else "not set"
        return DIFFERENT, it, code_has, f"standard {pct(std['value'], tgt)}, code {got}"
    return MISSING, None, "", ""


def _match_calculation(std: dict, items: list[dict]):
    for it in items:
        f = it["facts"]
        acts = f.get("actions", []) if f.get("kind") in ("statement", "condition") else []
        for a in acts:
            if a["type"] != "calculate" or not _all_parts(std["target"], a["target"]):
                continue
            uses_ok = all(any(_field_hit(u, [x]) for x in a.get("uses", [])) for u in std.get("uses", []))
            code_has = f'{a["target"]} = {a["expr"]}'
            if uses_ok:
                return MATCHED, it, code_has, ""
            return DIFFERENT, it, code_has, "calculation uses different inputs"
    return MISSING, None, "", ""


def _match_default(std: dict, items: list[dict]):
    for it in _tables(items):
        f = it["facts"]
        if _field_hit(std["table"]["field"], f.get("fields", [])):
            code_has = f'{f["table_id"]}: {", ".join(r["key"] for r in f["rows"])}'
            if f.get("has_default"):
                return MATCHED, it, code_has + " + default", ""
            return DIFFERENT, it, code_has + " (no default)", "no 'any other value' branch in code"
    return MISSING, None, "", ""


_USED: set = set()   # keys already matched in the current comparison

MATCHERS = {"condition": _match_condition, "lookup": _match_lookup,
            "calculation": _match_calculation, "default_case": _match_default}


def _expects(std: dict) -> str:
    t = std.get("type")
    if t == "condition":
        w = std["when"]
        then = "; ".join((f'{"/".join(_kw(e["set"]))} = {_fmt(e["value"])}' if "set" in e else
                          f'{"/".join(_kw(e["calculate"]))} × {e.get("factor", "")}') for e in std.get("then", []))
        return f'{_cond_text(w["field"], w["op"], w.get("value"))} → {then or "(any outcome)"}'
    if t == "lookup":
        return f'{"/".join(_kw(std["table"]["field"]))} = {std["key"]} → {"/".join(_kw(std["sets"]))} = {_fmt(std["value"])}'
    if t == "calculation":
        return f'{"-".join(_kw(std["target"]))} from {" × ".join("/".join(_kw(u)) for u in std.get("uses", []))}'
    if t == "default_case":
        return f'{"/".join(_kw(std["table"]["field"]))} table has a default for unknown values'
    return ""


# ---------------------------------------------------------------- compare
def compare(rules_json: dict, standard: dict) -> dict[str, Any]:
    items = _code_items(rules_json.get("programs", []))
    used = _USED
    used.clear()
    results = []
    for std in standard["rules"]:
        fn = MATCHERS.get(std.get("type"))
        if not fn:
            results.append({"std_id": std.get("id"), "name": std.get("name"), "status": "Not checked",
                            "detail": f"unknown type '{std.get('type')}'"})
            continue
        status, it, code_has, detail = fn(std, [i for i in items])
        if it is not None:
            used.add(it["key"])
            if it["facts"].get("kind") == "row":                # a table row matched: the table itself is covered
                used.add(it["key"].split(":")[0] + ":table")
        results.append({"std_id": std["id"], "category": std.get("category", ""), "name": std["name"],
                        "standard_rule": std.get("description", ""), "standard_expects": _expects(std),
                        "status": status, "code_has": code_has, "difference": detail,
                        "code_rule_id": it["rule_id"] if it else "", "program": it["program"] if it else "",
                        "source": it["lines"] if it else ""})
    additional = [{"code_rule_id": it["rule_id"], "program": it["program"], "type": it["type"],
                   "rule": it["rule"], "source": it["lines"], "status": ADDITIONAL}
                  for it in items if not it.get("internal") and it["key"] not in used]
    counts = {s: sum(1 for r in results if r["status"] == s) for s in (MATCHED, DIFFERENT, MISSING)}
    counts[ADDITIONAL] = len(additional)
    return {"standard": standard.get("standard", ""), "version": standard.get("version", ""),
            "owner": standard.get("owner", ""), "summary": counts,
            "comparison": results, "additional_in_code": additional}


# ---------------------------------------------------------------- render
def _cell(s) -> str:
    return (str(s) if s not in (None, "") else "—").replace("|", "\\|").replace("\n", " ")


_BADGE = {MATCHED: "✅ Matched", DIFFERENT: "🟠 Different", MISSING: "❌ Missing in code", ADDITIONAL: "➕ Additional in code"}


def render_md(cmp: dict) -> str:
    c = cmp["summary"]
    out = ["# Standards comparison\n",
           f"_Compared with **{cmp['standard']}** v{cmp['version']}"
           + (f" (owner: {cmp['owner']})" if cmp.get("owner") else "") + "._\n",
           "| ✅ Matched | 🟠 Different | ❌ Missing in code | ➕ Additional in code |",
           "|---|---|---|---|",
           f"| {c[MATCHED]} | {c[DIFFERENT]} | {c[MISSING]} | {c[ADDITIONAL]} |",
           "\n## Standard rules checked against the code\n",
           "| Std ID | Standard rule | Status | Standard expects | Code has | Difference | Code rule | Source |",
           "|---|---|---|---|---|---|---|---|"]
    order = {DIFFERENT: 0, MISSING: 1, MATCHED: 2}
    for r in sorted(cmp["comparison"], key=lambda r: (order.get(r["status"], 3), r["std_id"])):
        out.append(f"| {r['std_id']} | **{_cell(r['name'])}** — {_cell(r.get('standard_rule'))} | "
                   f"{_BADGE.get(r['status'], r['status'])} | {_cell(r.get('standard_expects'))} | "
                   f"{_cell(r.get('code_has'))} | {_cell(r.get('difference'))} | {_cell(r.get('code_rule_id'))} | "
                   f"{_cell(r.get('source'))} |")
    out += ["\n## Additional rules found in the code (not in the standard)\n",
            "_Behaviour the code performs that the standard does not describe. Each should be either added "
            "to the standard or confirmed as intentional._\n"]
    if cmp["additional_in_code"]:
        out += ["| Code rule | Type | Rule | Source |", "|---|---|---|---|"]
        for a in cmp["additional_in_code"]:
            out.append(f"| {a['code_rule_id']} | {a['type']} | {_cell(a['rule'])} | {_cell(a['source'])} |")
    else:
        out.append("_None — every extracted business rule is covered by the standard._")
    return "\n".join(out)


def write_xlsx(cmp: dict, path: str | Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    fills = {MATCHED: "D9F2E3", DIFFERENT: "FDEBD0", MISSING: "F8D7DA", ADDITIONAL: "DCE8F7"}
    head_fill, head_font = PatternFill("solid", fgColor="1F2A44"), Font(bold=True, color="FFFFFF")
    wrap = Alignment(wrap_text=True, vertical="top")

    def sheet(ws, headers, rows, widths, status_col=None):
        ws.append(headers)
        for c in ws[1]:
            c.fill, c.font, c.alignment = head_fill, head_font, Alignment(wrap_text=True, vertical="center")
        for row in rows:
            ws.append(row)
        for r in ws.iter_rows(min_row=2):
            for c in r:
                c.alignment = wrap
            if status_col is not None:
                st = r[status_col].value
                if st in fills:
                    r[status_col].fill = PatternFill("solid", fgColor=fills[st])
                    r[status_col].font = Font(bold=True)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    wb = Workbook()
    s = wb.active
    s.title = "Summary"
    s.append(["Standards comparison"])
    s["A1"].font = Font(bold=True, size=14)
    s.append(["Standard", cmp["standard"]])
    s.append(["Version", cmp["version"]])
    s.append(["Owner", cmp.get("owner", "")])
    s.append([])
    s.append(["Status", "Count"])
    for c in s[6]:
        c.fill, c.font = head_fill, head_font
    for st in (MATCHED, DIFFERENT, MISSING, ADDITIONAL):
        s.append([st, cmp["summary"][st]])
        s.cell(row=s.max_row, column=1).fill = PatternFill("solid", fgColor=fills[st])
    s.column_dimensions["A"].width, s.column_dimensions["B"].width = 24, 50

    sheet(wb.create_sheet("Comparison"),
          ["Std ID", "Category", "Standard rule", "Description", "Status", "Standard expects",
           "Code has", "Difference", "Code rule", "Program", "Source"],
          [[r["std_id"], r.get("category", ""), r["name"], r.get("standard_rule", ""), r["status"],
            r.get("standard_expects", ""), r.get("code_has", ""), r.get("difference", ""),
            r.get("code_rule_id", ""), r.get("program", ""), r.get("source", "")] for r in cmp["comparison"]],
          [9, 12, 26, 44, 18, 34, 34, 34, 10, 12, 22], status_col=4)
    sheet(wb.create_sheet("Additional in code"),
          ["Code rule", "Program", "Type", "Rule", "Source", "Status"],
          [[a["code_rule_id"], a["program"], a["type"], a["rule"], a["source"], a["status"]]
           for a in cmp["additional_in_code"]],
          [10, 12, 14, 70, 24, 20], status_col=5)
    wb.save(path)
